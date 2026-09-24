# -*- coding: utf-8 -*-

from PyQt5 import QtCore, QtGui, QtWidgets
import openpyxl as op
from err_dlg1 import ErrorDialog
import os
		
class Print_xlsx():
    def print_xl(self, records, mode1, wnd):
 
        path1 = u'D:\\PythonProj\\output\\'
        path2 = u'D:\\PythonProj\\django22\\pr1\\app1\\static\\xlsx\\'
        if (mode1 == 1):    
            wb = op.load_workbook(filename = path2 + 'userlist1.xlsm')
            ws = wb.active

            i = 5
            for elem in records:
                ws.cell(row = i, column = 1).value = elem[0]
                ws.cell(row = i, column = 2).value = elem[1]
                ws.cell(row = i, column = 3).value = elem[2]
                ws.cell(row = i, column = 4).value = elem[3]
                ws.cell(row = i, column = 5).value = elem[4]
                ws.cell(row = i, column = 6).value = elem[5]
                i = i + 1  
  
            fname = path1 + 'Список участников.xlsx'
            try:
                wb.save(fname)
            except Exception as e:
                self.parent.errmsg1 = 'Ошибка печати в файл: {}'.format(e)
                dialog = ErrorDialog(self.parent)
                dialog.show()

        elif (mode1 == 2):    
            wb = op.load_workbook(filename = path2 + 'anslist1.xlsm')
            ws = wb.active

            i = 5
            for elem in records:
                ws.cell(row = i, column = 1).value = elem[0]
                ws.cell(row = i, column = 2).value = elem[1]
                ws.cell(row = i, column = 3).value = elem[2]
                ws.cell(row = i, column = 4).value = elem[3]
                ws.cell(row = i, column = 5).value = elem[4]
                ws.cell(row = i, column = 6).value = elem[5]
                ws.cell(row = i, column = 7).value = elem[6]
                i = i + 1  
  
            fname = path1 + 'Список ответов.xlsx'
            try:
                wb.save(fname)
            except Exception as e:
                self.parent.errmsg1 = 'Ошибка печати в файл: {}'.format(e)
                dialog = ErrorDialog(self.parent)
                dialog.show()
			
